# -*- coding: utf-8 -*-
"""把 charts/ 的配图嵌进头条文稿的占位符位置（生成器自带函数的用法示例）

用法:
    powershell -File run.ps1 阿飞说球\02_工具\insert_charts_jleague_0912.py
"""
import os
import sys

BASE = r'D:\Hermes\football'
TOOL = os.path.join(BASE, '阿飞说球', '02_工具')
sys.path.insert(0, TOOL)
from afei_generator import insert_charts  # noqa: E402

ART = os.path.join(BASE, '阿飞说球', '04_赛事文章', '2026-09', '周六早场双场_日职')
CH = os.path.join(ART, 'charts')
DOCX = os.path.join(ART, '周六003+004日职：差1分净胜球差8个,钢巴盘口反向_发布导入版.docx')

PLAN = [
    # (锚点文本, 图片, 图注) —— 锚点用占位符行文本; 图片插在该行位置, 占位符随之删除
    ('【此处插入图1', os.path.join(CH, 'cmp1_双场模型vs市场偏差.png'),
     '图1 双场模型vs市场偏差对比'),
    ('【此处插入图2', os.path.join(CH, 'cmp2_双场交锋克星对比.png'),
     '图2 双场交锋对位对比'),
    ('【此处插入图3', os.path.join(CH, 'cmp3_四队攻防对比.png'),
     '图3 四队攻防数据对比(近10场)'),
]
if __name__ == '__main__':
    ok, missed = insert_charts(DOCX, PLAN)
    print(f'已嵌入 {ok} 张配图')
    if missed:
        print('未命中锚点:', missed)
    from docx import Document
    d = Document(DOCX)
    media = [str(x.partname) for x in d.part.package.parts if 'media' in str(x.partname)]
    left = [p.text for p in d.paragraphs if '此处插入图' in p.text]
    print('docx 内嵌图片:', len(media))
    print('剩余占位符:', len(left), left)
    print('图注:', [p.text for p in d.paragraphs if p.text.startswith('图')])
