# -*- coding: utf-8 -*-
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
import re
import os

doc = Document()

# 设置默认字体
style = doc.styles['Normal']
font = style.font
font.name = '微软雅黑'
font.size = Pt(11)

# 读取markdown文件
md_path = os.path.join("I:", "\u963f\u98de\u8bf4\u7403", "\U0001f4dd\u6587\u7ae0", "20260730", "\u6b27\u7f57\u5df4\u4e24\u573a\u524d\u77bb_\u672c\u83f2\u5361\u5723\u52a0\u4ed1_\u5b89\u5fb7\u83b1\u8d6b\u7279\u54c8\u9a6c\u6bd4_\u5934\u6761\u7248.md")
with open(md_path, "r", encoding="utf-8") as f:
    content = f.read()

lines = content.split('\n')

for line in lines:
    line = line.strip()
    
    # 跳过空行
    if not line:
        continue
    
    # 标题行（第一行）
    if line.startswith('\u5468\u56db003'):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = p.add_run(line)
        run.bold = True
        run.font.size = Pt(16)
        run.font.color.rgb = RGBColor(0x1a, 0x1a, 0x2e)
        continue
    
    # 标签行
    if line.startswith('#\u6b27\u7f57\u5df4'):
        p = doc.add_paragraph()
        run = p.add_run(line)
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0x99, 0x99, 0x99)
        continue
    
    # 分隔线
    if line == '---':
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run('\u2500' * 40)
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(0xdd, 0xdd, 0xdd)
        continue
    
    # 金句行
    if '\u6570\u636e\u4e0d\u4f1a\u8bf4\u8c0e' in line:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(line)
        run.italic = True
        run.font.size = Pt(12)
        run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
        continue
    
    # 配图关键词行（以数字开头）
    if re.match(r'^\d+-', line):
        p = doc.add_paragraph()
        run = p.add_run(line)
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(0x99, 0x99, 0x99)
        continue
    
    # 免责声明行
    if '\u884c\u6587\u4ed3\u4fc3' in line or '\u56fe\u4e00\u4e50' in line or '\u6570\u636e\u6765\u6e90' in line:
        p = doc.add_paragraph()
        run = p.add_run(line)
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0x99, 0x99, 0x99)
        continue
    
    # 普通段落
    p = doc.add_paragraph()
    run = p.add_run(line)
    run.font.size = Pt(11)

# 保存文档
output_path = os.path.join("I:", "\u963f\u98de\u8bf4\u7403", "\U0001f4dd\u6587\u7ae0", "20260730", "\u6b27\u7f57\u5df4\u4e24\u573a\u524d\u77bb_\u5934\u6761\u7248.docx")
doc.save(output_path)
print("Word document saved successfully")
