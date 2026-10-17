# -*- coding: utf-8 -*-
"""生成: 韩职双场002+003_平局大师撞克星_头条深度版.docx (R19 白底图规范)"""
import os
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt

plt.rcParams["font.sans-serif"] = ["Microsoft YaHei", "SimHei", "sans-serif"]
plt.rcParams["axes.unicode_minus"] = False

from docx import Document
from docx.shared import Pt, Cm, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml

BASE = r"I:\阿飞说球\_tmp_docx_out"
CHART_DIR = os.path.join(BASE, "charts")
os.makedirs(CHART_DIR, exist_ok=True)

C_MAIN = "4fc3f7"
C_DRAW = "ffd93d"
C_AWAY = "ff6b6b"
C_DARK = "1a237e"
C_GRAY = "9e9e9e"

def set_cn_font(run, size=10.5, bold=False, color=None, name="微软雅黑"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)

def shade_cell(cell, hexcolor):
    cell._tc.get_or_add_tcPr().append(parse_xml(r'<w:shd {} w:val="clear" w:fill="{}"/>'.format(nsdecls("w"), hexcolor)))

def para(doc, text, size=10.5, bold=False, color=None, align=None, space_after=6):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    r = p.add_run(text)
    set_cn_font(r, size=size, bold=bold, color=color)
    return p

def heading(doc, text, size=14, color=C_DARK, space_before=10):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    set_cn_font(r, size=size, bold=True, color=color)
    return p

def bar_chart(title, market, model, fname):
    cats = ["主胜", "平局", "客胜"]
    x = [0, 1, 2]
    w = 0.36
    fig, ax = plt.subplots(figsize=(7.2, 3.4), dpi=160)
    fig.patch.set_facecolor("white")
    ax.set_facecolor("white")
    colors = ["#" + c for c in (C_MAIN, C_DRAW, C_AWAY)]
    b1 = ax.bar([i - w / 2 for i in x], market, width=w, color=colors, label="市场定价", edgecolor="white")
    b2 = ax.bar([i + w / 2 for i in x], model, width=w, color="#37474f", alpha=0.78, label="后台模型", edgecolor="white")
    for bars in (b1, b2):
        for b in bars:
            ax.annotate(f"{b.get_height():.0f}%", (b.get_x() + b.get_width() / 2, b.get_height()),
                        ha="center", va="bottom", fontsize=8.5)
    ax.set_xticks(x)
    ax.set_xticklabels(cats, fontsize=10)
    ax.set_ylim(0, 75)
    ax.set_ylabel("概率 (%)", fontsize=9)
    ax.set_title(title, fontsize=11, color="#1a237e", pad=8)
    ax.legend(fontsize=8.5, loc="upper right", framealpha=0.9)
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    ax.grid(axis="y", linestyle="--", alpha=0.35)
    fig.tight_layout()
    fig.savefig(fname, facecolor="white", bbox_inches="tight")
    plt.close(fig)

ch1 = os.path.join(CHART_DIR, "bar_m1.png")
ch2 = os.path.join(CHART_DIR, "bar_m2.png")
bar_chart("济州SK vs 大田市民 · 胜平负概率", [25, 27, 48], [32, 24, 44], ch1)
bar_chart("仁川联 vs 全北现代 · 胜平负概率", [30, 30, 40], [30, 33, 37], ch2)

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.left_margin = sec.right_margin = Cm(2.2)
sec.top_margin = sec.bottom_margin = Cm(2.0)

style = doc.styles["Normal"]
style.font.name = "微软雅黑"
style.element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
style.font.size = Pt(10.5)

p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("韩职双场002+003：平局大师撞克星，客让1球全落空")
set_cn_font(r, size=17, bold=True, color=C_DARK)
p.paragraph_format.space_after = Pt(4)

para(doc, "#济州SK  #大田市民  #仁川联  #全北现代  #韩职  #足球前瞻  #数据分析", size=9, color=C_GRAY,
     align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10)

para(doc, "很多人看韩职早场，只看排名：大田第9、全北第3，一个状态火，一个豪门底，以为客场稳了。但后台模型跑完这两场，发现一个共性：两个客场热门，客让1球都很难打出来。济州是主场平局大师（赛季4胜6平5负、近10个主场平了5场），大田还伤了前锋；全北更离谱，客场13场只赢4场、平了6场，仁川还是它的克星（近3次交手2胜1平，40天前刚1:0赢过）。", size=10.5)
para(doc, "结论先行：大田净胜2球双共识只有23%，全北更惨只有15-17%。客让1球两场都是坑，受让一侧才是保护层；而两场的底色都是平局（济州主场40%平率、全北客场46%平率）。", size=10.5, bold=True, color=C_DARK)

heading(doc, "第一场：济州SK vs 大田市民 — 主场平局大师，客场猎手让1球是陷阱")
para(doc, "很多人看到大田近10场6胜2平2负、刚赢蔚山，觉得客场稳了。但济州的主场平局体质是韩职出了名的：赛季主场15场只赢4场、平了6场，近10个主场平了5场（50%）。", size=10.5)
para(doc, "⚔️ 战意：济州第5（35分）争冠组边缘，主场不能再丢分；大田第9（32分）无保级压力，09-15 还有亚冠精英首秀的激励，大田士气在线，但也不是生死战。", size=10.5)
doc.add_picture(ch1, width=Inches(5.9))
doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

tbl = doc.add_table(rows=5, cols=4)
tbl.style = "Table Grid"; tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
hdr = ["结果", "市场定价", "后台模型", "偏差"]
for j, h in enumerate(hdr):
    c = tbl.rows[0].cells[j]; c.text = ""
    r = c.paragraphs[0].add_run(h); set_cn_font(r, size=9.5, bold=True, color="ffffff")
    shade_cell(c, C_DARK)
rows = [
    ("济州胜", "25%", "32%", "+6.6pp", C_MAIN),
    ("平局", "27%", "24%", "-2.8pp", C_DRAW),
    ("大田胜", "48%", "44%", "-3.8pp", C_AWAY),
    ("大2.5", "≈51%", "63%", "+12pp", C_MAIN),
]
for i, (a, b, c_, d, tint) in enumerate(rows, start=1):
    for j, v in enumerate([a, b, c_, d]):
        cell = tbl.rows[i].cells[j]; cell.text = ""
        r = cell.paragraphs[0].add_run(v); set_cn_font(r, size=9.5, bold=(j == 3))
        if j < 3:
            shade_cell(cell, tint)
para(doc, "", size=4, space_after=2)
para(doc, "💰 信号：市场的 48% 客热有大田状态背书（近5联赛3胜1平1负），但客让1球是陷阱：大田要净胜2球，双共识只有23-24%，何况前锋迪奥戈还伤停、中场严元相也伤。比分市场第一重心是 1:1（全场最低价），平局小比分才是市场真正的主线；大小球模型63% vs 市场51%有分歧，但H2H场均2.2球的小球惯性站在市场侧，不追大。", size=10.5)
para(doc, "🌍 外部视角：综合数据倾向：韩媒常规叙事「济州主场闷平体质 vs 大田状态火热」，两队近3次交锋2平1负（大田视角），支持小比分拉锯。", size=10.5)
para(doc, "高概率比分：首选 1-1，次选 1-2，留意 0-1。", size=10.5, bold=True, color=C_DARK)

heading(doc, "第二场：仁川联 vs 全北现代 — 客场平局大师撞上自己的克星")
para(doc, "很多人看到全北排名第3、韩职豪门，以为客场打中游手拿把攥。但全北本赛季客场13场只赢4场、平了6场（胜率31%、平率46%），它是韩职最典型的客场平局大师；而仁川近3次交手2胜1平、40天前刚在主场1:0赢过它，是公认的克星。", size=10.5)
para(doc, "⚔️ 战意：全北第3争冠组，09-16 还有亚冠精英 vs 柏太阳神，激励在，但临近洲际也可能留力；仁川第8中游无压力，主场对全北有心理底气。", size=10.5)
doc.add_picture(ch2, width=Inches(5.9))
doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

tbl2 = doc.add_table(rows=5, cols=4)
tbl2.style = "Table Grid"; tbl2.alignment = WD_TABLE_ALIGNMENT.CENTER
for j, h in enumerate(hdr):
    c = tbl2.rows[0].cells[j]; c.text = ""
    r = c.paragraphs[0].add_run(h); set_cn_font(r, size=9.5, bold=True, color="ffffff")
    shade_cell(c, C_DARK)
rows2 = [
    ("仁川胜", "30%", "30%", "0pp", C_MAIN),
    ("平局", "30%", "33%", "+3.2pp", C_DRAW),
    ("全北胜", "40%", "37%", "-3.2pp", C_AWAY),
    ("大2.5", "≈55%", "31%", "-24pp", C_MAIN),
]
for i, (a, b, c_, d, tint) in enumerate(rows2, start=1):
    for j, v in enumerate([a, b, c_, d]):
        cell = tbl2.rows[i].cells[j]; cell.text = ""
        r = cell.paragraphs[0].add_run(v); set_cn_font(r, size=9.5, bold=(j == 3))
        if j < 3:
            shade_cell(cell, tint)
para(doc, "", size=4, space_after=2)
para(doc, "💰 信号：全北客让1球要净胜2球，双共识只有15-17%，是两场里最难打出的一个；模型平局32.7%上浮（市场29.5%）：全北客场46%平率 + H2H近10场4平 + 仁川主场对全北近5次2胜2平1负，平局是本场最浓的信号。大小球分歧比第一场更尖锐：模型31% vs 市场55%，H2H场均2.0球站在模型侧，不追大。", size=10.5)
para(doc, "🌍 外部视角：综合数据倾向：韩媒叙事「全北客场平局体质 vs 仁川专克全北」，全北视角近3次交锋1平2负，拉锯闷局是主线。", size=10.5)
para(doc, "高概率比分：首选 0-0，次选 0-1，留意 1-1。", size=10.5, bold=True, color=C_DARK)

heading(doc, "最终判断：韩职双场矩阵")
mt = doc.add_table(rows=3, cols=4)
mt.style = "Table Grid"; mt.alignment = WD_TABLE_ALIGNMENT.CENTER
mhdr = ["场次", "后台模型", "高概率比分", "核心信号"]
for j, h in enumerate(mhdr):
    c = mt.rows[0].cells[j]; c.text = ""
    r = c.paragraphs[0].add_run(h); set_cn_font(r, size=9.5, bold=True, color="ffffff")
    shade_cell(c, C_DARK)
mrows = [
    ("济州 vs 大田", "客44% vs 市场48%", "1-1 / 1-2 / 0-1", "客让1球打穿仅23%（大田锋线伤停）；主+1 55.9%"),
    ("仁川 vs 全北", "客37% vs 市场40%", "0-0 / 0-1 / 1-1", "平局32.7%；客让1球打穿仅15-17%；主+1 62.7%"),
]
for i, row in enumerate(mrows, start=1):
    for j, v in enumerate(row):
        c = mt.rows[i].cells[j]; c.text = ""
        r = c.paragraphs[0].add_run(v); set_cn_font(r, size=9.5, bold=(j == 3))
para(doc, "", size=4, space_after=2)
para(doc, "两场共同规律：客场热门都不好赢两球（净胜2球双共识15-24%），受让一侧（济州+1 / 仁川+1）是双场共识的保护层；而平局是共同的底色（济州主场40%平率、全北客场46%平率）。说白了，这两场的钱不在\"客队赢\"，在\"客队赢不了两球\"和\"闷平\"。", size=10.5, bold=True)

heading(doc, "互动钩子", size=12)
para(doc, "韩职双场拆完了。大田状态火热，你更信它客场能赢两球，还是济州主场继续闷？全北客场平了6场，它这场终于开窍还是继续闷？评论区留下你的预测比分，赛后我们用实战赛果复盘模型精度。", size=10.5)
para(doc, "难说。真说不好。", size=10.5, bold=True)
para(doc, "数据来源：懂球帝、ESPN、各队官网", size=8.5, color=C_GRAY)
para(doc, "注：本文基于两队技战术、历史数据与赛事背景的客观模型推演，仅供足球运动战术探讨，请理性看待比赛。", size=8.5, color=C_GRAY)

out = os.path.join(BASE, "韩职双场002+003_平局大师撞克星_头条深度版.docx")
doc.save(out)
print("SAVED:", out)
