# python-docx 生成模板

> 用于头条号文章导出，表格在Word中正常渲染。

```python
from docx import Document
from docx.shared import Pt

doc = Document()

def add_body(text, bold=False, size=12):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.size = Pt(size)
    if bold: run.bold = True
    return p

def add_title(text, size=13):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.bold = True
    return p

# 标题
doc.add_heading('标题', level=1)

# 信息头
info = doc.add_paragraph()
r = info.add_run('赛事信息\n时间 · 场地 · 轮次')
r.font.size = Pt(11)

# 正文
add_body('段落内容...')

# 表格
table = doc.add_table(rows=7, cols=4, style='Table Grid')
headers = ['维度', '主队', '客队', '优势方']
data = [
    ('场均进球', 'X.X', 'X.X', '主/客'),
    ('场均失球', 'X.X', 'X.X', '主/客'),
    # ... 更多行
]
for i, h in enumerate(headers):
    table.cell(0, i).text = h
    for p in table.cell(0, i).paragraphs:
        for r in p.runs: r.bold = True
for i, row in enumerate(data):
    for j, val in enumerate(row):
        table.cell(i+1, j).text = val

# 签名
p = doc.add_paragraph()
r = p.add_run('数据不会说谎，但人会。——阿飞·赛博聊球')
r.bold = True

# 免责
add_body('本文仅为赛事数据分析，不构成任何投注建议。竞猜有风险，请理性参与。')
add_body('数据来源：竞彩网、咪咕视频、懂球帝')

doc.save('output.docx')
```

## 使用要点
- `style='Table Grid'` 给表格加网格线
- 表头行加粗：遍历cell的paragraphs的runs设bold
- 段落间用`doc.add_paragraph()`空行分隔
- 保存路径用项目目录下的日期子文件夹
