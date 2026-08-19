# -*- coding: utf-8 -*-
"""
周四002 巴列卡诺vs阿拉维斯 头条版 docx 生成 (R19: 白底图/独立交付)
输出: 📝文章/20260821/周四002_巴列卡诺vs阿拉维斯_头条版.docx
"""
import sys, os
sys.stdout.reconfigure(encoding='utf-8')

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

CHART_DIR = r'I:\阿飞说球\database\charts_20260821'
OUT = r'I:\阿飞说球\📝文章\20260821\周四002_巴列卡诺vs阿拉维斯_头条版.docx'
C_PRIMARY = RGBColor(0x0e, 0x66, 0x55)
C_DARK = RGBColor(0x33, 0x33, 0x33)
C_GRAY = RGBColor(0x88, 0x88, 0x88)

TITLE = '周四002巴列卡诺VS阿拉维斯：3-0大胜撞上四连杀克星'
TAGS = '#巴列卡诺 #阿拉维斯 #西甲 #足球前瞻 #数据分析 #赛博聊球 #西甲第2轮 #足球推荐 #体育头条 #足球分析'

BODY = [
    ('p', '阿拉维斯这队，首轮3-0把赫塔费收拾得明明白白，现在排第二。巴列卡诺呢，首轮客场1-2输给塞维利亚，一分没拿，排到第18去了。光看这俩排名，谁都会觉得客队更有底气。'),
    ('p', '但翻翻这两队的历史账本，事情就不太一样了。近10次交手，巴列卡诺赢了7场，输了3场，关键是一局平局都没有。主场打阿拉维斯，连着4次全赢，比分就俩：1-0、2-0。你说阿拉维斯新赛季状态热，可对面这队，天生就是它的克星。'),
    ('h', '先掰扯数据'),
    ('p', '巴列卡诺最近10个主场，6胜3平1负，主场这关是它吃饭的家伙。场均进1.2丢1.3，不算好看，但能赢球。阿拉维斯最近10场5胜3平2负，账面更漂亮，可你要拆开看客场：4胜3平3负，一般般。场均进1.5丢0.9，这数字是含了首轮那场3-0的水分的。大胜是主场踢出来的，客场啥水平，还真不好说。'),
    ('p', '伤停这边，阿拉维斯主力后卫F. Garcés停赛，防空和二点球保护直接打折。巴列卡诺暂无伤停，但进球全靠一个人扛，Álvaro García，全队100%的球都是他进的，他被盯死了，这队就哑火。'),
    ('h', '打法对位，历史账本对得上'),
    ('p', '巴列卡诺主场那套高位逼抢，专门收拾后场出球不利索的队。阿拉维斯客场本来就爱缩着打反击，现在后卫还停了一个，更不敢跟主队对攻了。这套打法对位，跟近10次交锋的结果完全对得上。'),
    ('p', '这里有个信号挺扎眼：两队近10次交手，0场平局。要么你赢，要么我赢，就是没有平手。可机构给平局定了31.2%的概率，拿什么撑起来的？交锋记录里可找不到依据。'),
    ('h', '题材热归热，数据端没跟着热'),
    ('p', '客队3-0大胜，题材这么热，数据端反应了啥？客胜从3.70降到3.58，就意思了一下；主胜一直钉在2.01，市场隐含44.1%，压根没破2.0。机构拿客队的热度当幌子，主胜的筹码该收还是收。'),
    ('p', '再看让球。主队让一球，净胜两球才19.3%的可能，机构54.7%押客队+1不败。这不是没道理，历史交锋主队全赢，但都是1球、2球的小胜，就没打过3球。总进球2球定价最低（3.20），机构明摆着看小，跟两队交锋场均1.6球对上了。'),
    ('h', '📊 数据速览'),
    ('p', '巴列卡诺(主)：第18名0分，近10场4胜2平4负，同主客6胜3平1负，场均进1.2/失1.3，近10场交锋7胜0平3负，主场交锋近4次全胜，暂无伤停'),
    ('p', '阿拉维斯(客)：第2名3分，近10场5胜3平2负，同主客4胜3平3负，场均进1.5/失0.9，交锋3胜0平7负，F. Garcés(后卫)停赛'),
    ('c', 'bar_rayo'),
    ('h', '🧮 模型'),
    ('p', '后台模型跑完，λ主1.45对λ客0.95，巴列卡诺胜率48.7%，平局26.6%，阿拉维斯24.7%。模型的主胜比机构定价高了4.6个点，平局低了4.6个点。分歧就出在平局上，模型觉得机构给的平局概率虚高。客胜两边倒是没啥差距，没便宜可捡。'),
    ('c', 'radar_rayo'),
    ('h', '🌍 外媒'),
    ('p', '外媒这边，方向倒是出奇一致，全看主队不败。Sports Mole说主队小胜，Dailysports的模型给主队45-50%，Meczyki对让一球能不能净胜两球不太放心。分歧只在赢几个。有外媒喊大胜，我看悬，历史交锋全是1-2球的小胜，大胜是外媒的乐观，不是数据的底牌。'),
    ('h', '💰 信号'),
    ('p', '反正我的判断：克星压制、主场优势、对手后卫停赛，三条叠一起，主队不败概率能到75%。但主胜的定价优势只有4.6个点，中规中矩，不是那种闭着眼都能上的强信号。巴列卡诺首轮刚输球，进攻又单点依赖，这场的底气，一大半是历史交锋给的，一小半才是自己踢出来的。老实说，我心里也没底，上赛季我也信过一次"克星"，被狠狠打过脸。'),
    ('p', '让球这边更直白。主队让一球，要净胜两球才算赢下，可历史4个主场就2场净胜2球，另外2场只赢1个。让一球正好卡在走盘的边儿上，占不到便宜。你要是问我让球玩不玩，我的态度就俩字：算了。'),
    ('h', '🎬 剧本'),
    ('p', '剧本A（主队控制小胜，约45%）：巴列卡诺主场压着打，阿拉维斯缩着守还缺人，1-0或者2-0，这是两队历史交锋里最常见的剧情。这种走法，主胜加小球。'),
    ('p', '剧本B（闷平或客队偷袭，约30%）：主队破密集防守的老毛病犯了，前60分钟打不开局面，0-0或者1-1；要是被阿拉维斯反击偷一个，就是0-1。这种走法，平局加小球，冷门方向留意。'),
    ('p', '剧本C（客队爆冷，约25%）：阿拉维斯把首轮的状态带到客场，M.Díaz反击得手，1-2。概率不高，但人家首轮3-0是真的踢出来的，别当不存在。'),
    ('p', '写到这里我自己都想笑：三个剧本加起来，我还是在说"主队大概率赢但赢得不轻松"。看球这事就这样，分析来分析去，最后都是废话文学。算了，直接上结论。'),
    ('h', '🎯 推荐'),
    ('c', 'score_rayo'),
    ('p', '推荐比分：首选 1-0（14.1%）→ 次选 2-0（10.2%）→ 留意 1-1（9.9%）。冷门方向 0-1（7.4%）。四重落脚：方向看主队保平争胜（不败75%）；门槛看主队让-1难净胜两球，客队+1有支撑；进球看胶着小比分1-2球；比分首选1-0，次选2-0，防范1-1。'),
    ('h', '互动'),
    ('p_bold', '近10次交锋一场平局都没有，机构却把平局定到31.2%。你说这规律该终结了，还是机构这次看走眼了？评论区留个比分，赛后回来验证。我先押1-0，克星主场连赢4次，翻盘的剧本写不出来。'),
    ('p_gray', '数据来源：竞彩官方及新球体育（titan007）。'),
    ('p_gray', '纯属个人瞎琢磨，理性看球，远离赌博。'),
    ('p_gray', '注：本文为基于两队技战术、历史数据与赛事背景的客观模型推演，仅供足球运动战术探讨与交流，请理性看待比赛。'),
]


def set_font(run, size=11, bold=False, color=C_DARK, name='微软雅黑'):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = name
    run._element.rPr.rFonts.set(qn('w:eastAsia'), name)


def build_doc():
    doc = Document()
    style = doc.styles['Normal']
    style.font.name = '微软雅黑'
    style.font.size = Pt(11)
    style.element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(TITLE)
    set_font(run, size=16, bold=True, color=C_PRIMARY)

    p = doc.add_paragraph()
    run = p.add_run(TAGS)
    set_font(run, size=9, color=C_GRAY)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run('数据不会说谎，但人会。——阿飞·赛博聊球')
    set_font(run, size=10, color=C_GRAY)

    for item in BODY:
        kind = item[0]
        if kind == 'p':
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run(item[1])
            set_font(run, size=11)
        elif kind == 'p_bold':
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run(item[1])
            set_font(run, size=11, bold=True)
        elif kind == 'p_gray':
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run(item[1])
            set_font(run, size=9, color=C_GRAY)
        elif kind == 'h':
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(item[1])
            set_font(run, size=13, bold=True, color=C_PRIMARY)
        elif kind == 'c':
            fname = item[1] + '.png'
            path = os.path.join(CHART_DIR, fname)
            if os.path.exists(path):
                doc.add_picture(path, width=Inches(5.5))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p = doc.add_paragraph()
                run = p.add_run('[图表缺失: %s]' % fname)
                set_font(run, size=10, color=C_GRAY)

    doc.save(OUT)
    print('DOCX saved: %s (%d KB)' % (OUT, os.path.getsize(OUT) // 1024))


build_doc()
print('完成')
