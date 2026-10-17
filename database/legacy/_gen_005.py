# -*- coding: utf-8 -*-
# 周六005 头条版配图 + docx 生成 (临时 runner, 用完即删)
import os, re
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml.ns import qn

plt.rcParams["font.sans-serif"] = ["Microsoft YaHei", "SimHei"]
plt.rcParams["axes.unicode_minus"] = False

OUT = r"I:\阿飞说球\database"
F1 = os.path.join(OUT, "_b36_fig1.png")
F2 = os.path.join(OUT, "_b36_fig2.png")

BLUE, YELLOW, RED = "#4fc3f7", "#ffd93d", "#ff6b6b"
GRAY, TEAL = "#888780", "#1d9e75"

# ---------------- 图1: 胜平负 市场 vs 后台模型 ----------------
labels = ["热刺主胜", "平局", "维拉客胜"]
market = [48.66, 25.67, 25.67]
modelv = [42.3, 30.0, 27.6]
x = range(len(labels))
w = 0.36
fig, ax = plt.subplots(figsize=(7.2, 3.6), dpi=160)
fig.patch.set_facecolor("#ffffff"); ax.set_facecolor("#ffffff")
b1 = ax.bar([i - w/2 for i in x], market, w, label="市场定价", color=BLUE, edgecolor="#185fa5", linewidth=0.6)
b2 = ax.bar([i + w/2 for i in x], modelv, w, label="后台模型(折中口径)", color=RED, edgecolor="#a32d2d", linewidth=0.6)
for b in list(b1) + list(b2):
    ax.text(b.get_x() + b.get_width()/2, b.get_height() + 1.0,
            f"{b.get_height():.1f}%", ha="center", va="bottom", fontsize=9, color="#2C2C2A")
ax.set_xticks(list(x)); ax.set_xticklabels(labels, fontsize=11)
ax.set_ylabel("概率", fontsize=11)
ax.set_title("周六005 胜平负：市场定价 vs 后台模型", fontsize=13, pad=12)
ax.set_ylim(0, 58)
ax.legend(frameon=False, fontsize=10, loc="upper right")
for s in ("top", "right"): ax.spines[s].set_visible(False)
for s in ("left", "bottom"): ax.spines[s].set_color("#D3D1C7")
ax.grid(axis="y", color="#F1EFE8", linewidth=0.8)
fig.tight_layout(); fig.savefig(F1, facecolor="#ffffff"); plt.close(fig)

# ---------------- 图2: 大球概率 口径分歧 ----------------
names = ["联赛状态口径", "折中口径\n(叠加伤停)", "历史交锋口径", "市场定价"]
vals = [19.2, 37.7, 59.9, 56.81]
cols = [BLUE, TEAL, YELLOW, GRAY]
fig, ax = plt.subplots(figsize=(7.2, 3.6), dpi=160)
fig.patch.set_facecolor("#ffffff"); ax.set_facecolor("#ffffff")
b = ax.bar(names, vals, 0.55, color=cols, edgecolor="#5F5E5A", linewidth=0.6)
for r, v in zip(b, vals):
    ax.text(r.get_x() + r.get_width()/2, v + 1.5, f"{v}%", ha="center", va="bottom",
            fontsize=10, color="#2C2C2A", fontweight="bold")
ax.set_ylabel("三球及以上概率", fontsize=11)
ax.set_title("大小球：三个口径差了三倍，市场站在历史惯性一边", fontsize=13, pad=12)
ax.set_ylim(0, 72)
for s in ("top", "right"): ax.spines[s].set_visible(False)
for s in ("left", "bottom"): ax.spines[s].set_color("#D3D1C7")
ax.grid(axis="y", color="#F1EFE8", linewidth=0.8)
fig.tight_layout(); fig.savefig(F2, facecolor="#ffffff"); plt.close(fig)

# ---------------- md -> docx ----------------
SRC = r"I:\阿飞说球\📝文章\20260919\周六005_热刺vs维拉_四轮零进球凭什么净胜两球_赛事前瞻.md"
DST = r"I:\阿飞说球\20260919_周六005_热刺vs维拉_四轮零进球凭什么净胜两球_赛事前瞻.docx"
md = open(SRC, encoding="utf-8").read()
lines = md.split("\n")

doc = Document()
st = doc.styles["Normal"]
st.font.name = "Microsoft YaHei"; st.font.size = Pt(11)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")

def add_runs(p, text):
    """处理 **加粗**"""
    for part in re.split(r"(\*\*[^*]+\*\*)", text):
        if not part: continue
        if part.startswith("**") and part.endswith("**"):
            r = p.add_run(part[2:-2]); r.bold = True
        else:
            p.add_run(part)

IMG_AFTER = {"## 市场给的剧本": F1, "## 后台模型的选择": F2}

i = 0
while i < len(lines):
    ln = lines[i].rstrip()
    if not ln.strip():
        i += 1; continue
    if ln.startswith("# "):
        h = doc.add_heading(ln[2:].strip(), level=1)
        for r in h.runs: r.font.color.rgb = RGBColor(0x2C, 0x2C, 0x2A)
        i += 1; continue
    if ln.startswith("## "):
        h = doc.add_heading(ln[3:].strip(), level=2)
        for r in h.runs: r.font.color.rgb = RGBColor(0x18, 0x5F, 0xA5)
        if ln.strip() in IMG_AFTER:
            doc.add_picture(IMG_AFTER[ln.strip()], width=Inches(5.6))
            doc.paragraphs[-1].alignment = 1
        i += 1; continue
    if ln.strip() == "---":
        i += 1; continue
    if ln.startswith("> "):
        p = doc.add_paragraph(); add_runs(p, ln[2:]); p.paragraph_format.left_indent = Inches(0.3)
        i += 1; continue
    p = doc.add_paragraph(); add_runs(p, ln)
    i += 1

doc.save(DST)
print("DOCX OK ->", DST, os.path.getsize(DST), "bytes")
print("FIG1", F1, os.path.getsize(F1), "| FIG2", F2, os.path.getsize(F2))
