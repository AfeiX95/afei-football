# -*- coding: utf-8 -*-
import sys, os
sys.stdout.reconfigure(encoding="utf-8", errors="replace")
from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

BASE = r"I:\阿飞说球\📝文章\20260807"
CHARTS = os.path.join(BASE, "charts", "mp")
OUTPUT = os.path.join(BASE, "日职揭幕两场赛事前瞻_干货版.docx")

doc = Document()
style = doc.styles["Normal"]
style.font.name = "Microsoft YaHei"
style.font.size = Pt(10.5)
style.element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
style.paragraph_format.space_after = Pt(4)
style.paragraph_format.line_spacing = 1.25
for s in doc.sections:
    s.top_margin = Cm(1.8); s.bottom_margin = Cm(1.8)
    s.left_margin = Cm(2.2); s.right_margin = Cm(2.2)

def heading(text, level=1):
    h = doc.add_heading(text, level=level)
    for r in h.runs:
        r.font.color.rgb = RGBColor(0x1A, 0x1A, 0x2E)
    return h

def para(text, bold=False, size=10.5, color=None, align=None, italic=False):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.name = "Microsoft YaHei"; r.font.size = Pt(size)
    r.element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    if bold: r.bold = True
    if italic: r.italic = True
    if color: r.font.color.rgb = color
    if align: p.alignment = align
    return p

def add_image(path, width=Inches(5.5)):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(); r.add_picture(path, width=width)

def caption(text):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text); r.font.size = Pt(8); r.font.color.rgb = RGBColor(0x99, 0x99, 0x99)

def tbl(headers, rows):
    t = doc.add_table(rows=1+len(rows), cols=len(headers))
    t.style = "Light Grid Accent 1"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]; c.text = h
        for p in c.paragraphs:
            for r in p.runs:
                r.bold = True; r.font.size = Pt(9.5)
    for ri, row in enumerate(rows, start=1):
        for ci, val in enumerate(row):
            c = t.rows[ri].cells[ci]; c.text = val
            for p in c.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(9.5)
    return t

def img(name, cap):
    add_image(os.path.join(CHARTS, name), width=Inches(5.5))
    caption(cap)

# ===== 封面与金句 =====
para("阿飞·赛博聊球 ｜ 日职联 2026 揭幕战 ｜ 双场干货版", size=9, color=RGBColor(0x88,0x88,0x88), align=WD_ALIGN_PARAGRAPH.CENTER)
para("数据不会说谎，但人会。- 阿飞·赛博聊球", size=11, italic=True, color=RGBColor(0x33,0x33,0x33), align=WD_ALIGN_PARAGRAPH.CENTER)

# ===== 主标题与开篇 =====
heading("日职揭幕两场前瞻：数据在打架，市场只信了一半", level=1)
para("第一场，市场把鹿岛客胜定价到50.6%，后台模型只给32.6%。18个百分点的偏差，是开球前最值钱的一条信息。")
para("第二场，大阪钢巴排第4，浦和红钻排第17，账面像碾压局。翻开庭账，钢巴主场对浦和近10场只赢1场、输掉7场，4天后还要飞去客场踢亚冠附加赛。排名和现实，总有一个在说谎。")
para("两场同一个剧本：数据在打架，市场只信了一半。")

heading("核心结论", level=2)
para("横滨水手对鹿岛鹿角：后台模型主胜37.4%、平局30.0%、客胜32.6%；市场把客胜定价到50.6%，虚高约18个百分点。鹿岛客场近况差、双主力伤停，平局和横滨受让是价值点。", bold=True)
para("大阪钢巴对浦和红钻：账面主胜43.5%，经交锋与赛程修正后掉到32%，浦和反而35%略占先。克星效应加亚冠分心，别只看排名。", bold=True)

# ===== 第一场 =====
heading("第一场｜鹿岛客胜被高估18个百分点，横滨受让值不值得信？（18:25）", level=1)
tbl(["对阵", "横滨水手（主）vs 鹿岛鹿角"], [
    ["联赛", "日职联 2026 第1轮"],
    ["场地", "日产体育场 30°C 间中有云"],
    ["伤停", "横滨：7号前锋远野大弥（伤病）｜ 鹿岛：2号后卫安西幸辉（伤病）、14号中场樋口雄太（腓骨骨折）"],
])

heading("【数据偏差池】", level=2)
tbl(["方向", "后台模型", "市场定价", "偏差"], [
    ["横滨胜", "37.4%", "22.7%（3.90）", "+14.7pp"],
    ["平局", "30.0%", "26.7%（3.32）", "+3.3pp"],
    ["鹿岛胜", "32.6%", "50.6%（1.75）", "-18.0pp"],
])
para("偏差结论：鹿岛客胜被高估18个百分点，横滨主胜被低估约15个百分点。市场在买名气，模型在看近况。")

heading("【战术胜负手】", level=2)
para("鹿岛客场反击效率撑不起身价：近10个客场只赢4场、场均进球1.0，横滨主场对鹿岛近5场4胜1平。破局点在横滨的高位压迫，不在鹿岛的防反。")

heading("【伤停链条】", level=2)
para("横滨：7号前锋远野大弥缺阵，锋线少一个支点，但主场破门点不止一个，体系影响有限。")
para("鹿岛：2号后卫安西幸辉（伤病）和14号中场樋口雄太（腓骨骨折）同时缺阵。后卫到中场的出球链要临时拼凑，恰好撞上横滨的主场压迫，这是本场最大的战术软肋。")

heading("【战意与赛程】", level=2)
para("揭幕战，双方都要开门红，不存在留力动机；两队休整天数相同，赛程公平。30°C高温是唯一的环境变量，节奏偏慢，下半场体能决定比赛走向。")

heading("【数据验证】", level=2)
para("模型λ估算1.38对1.33，账面五五开；期望总进球2.71，大2.5概率51%，首选中低比分。")
para("数据平台特征分析：鹿岛客场场均进球1.0、近3个联赛客场不胜，和50.6%的客胜定价明显错位；横滨虽然也缺主力前锋，但主场进攻体系完整，破门点不止一个。")

img("m1_radar.png", "图1｜六维战力对比：横滨交锋占优，鹿岛客场疲软")
img("m1_bar.png", "图2｜胜平负概率：模型 vs 市场定价（鹿岛客胜被高估18pp）")
img("m1_pie.png", "图3｜修正后概率：平局第一")

heading("【推荐】", level=2)
para("方向：平局第一（修正后约35%），横滨主场不败约67%，不追鹿岛客胜。", bold=True)
para("比分：首选1-1，次选1-0，留意0-1。")
para("让球：鹿岛客让半球附近，看好横滨受让不败。")
para("大小：2-3球区间，大2.5谨慎偏小，高温揭幕战节奏慢，双方近期小球占比高。")
img("m1_signals.png", "图4｜信号图：鹿岛客场成色不足")

# ===== 第二场 =====
heading("第二场｜钢巴第4名却10战浦和7负，亚冠分心还信不信？（18:30）", level=1)
tbl(["对阵", "大阪钢巴（主）vs 浦和红钻"], [
    ["联赛", "日职联 2026 第1轮"],
    ["赛程", "钢巴8-11客场踢亚冠精英附加赛（对江原FC），4天后"],
    ["伤停", "双方暂无官方伤停数据"],
])

heading("【数据偏差池】", level=2)
tbl(["方向", "基础模型", "综合修正", "说明"], [
    ["钢巴胜", "43.5%", "32%", "交锋1胜7负+亚冠分心，下调11pp"],
    ["平局", "28.2%", "33%", "日职揭幕战平局概率上浮"],
    ["浦和胜", "28.4%", "35%", "客场克星+防守顶级，上调7pp"],
])
para("偏差结论：账面主胜被高估约11个百分点。克星效应加赛程分心，把排名优势彻底抹平，浦和反超为第一概率。")

heading("【战术胜负手】", level=2)
para("浦和低位防反对钢巴边路强攻：近5次联赛交手3次零封，浦和近10场场均失球0.7。钢巴场均2.1球的火力，专打浦和打不动。")

heading("【轮换与战意】", level=2)
para("官方伤停：双方暂无数据。真正的变量是轮换：钢巴4天后要踢亚冠精英附加赛，客场打江原FC，揭幕战大概率收着踢、保人不保分；浦和下一场要等一周，这场是全力抢分的最佳窗口。战意差一截。")

heading("【数据验证】", level=2)
para("浦和近10场5胜3平2负、场均失球0.7，近6场5场不败；钢巴近10场5胜2平3负、场均进球2.1，但赢的对手多是中下游，含金量要打折。")
para("数据平台特征分析：浦和防守评分明显高于钢巴，客场遇强队反击效率不差；交锋近5场4场不超过2球，小球基因扎实。")

img("m2_radar.png", "图5｜六维战力对比：浦和防守+交锋双压制")
img("m2_bar.png", "图6｜胜平负概率：基础模型 vs 综合修正（H2H+分心）")
img("m2_pie.png", "图7｜修正后概率：浦和略占优")

heading("【推荐】", level=2)
para("方向：浦和客胜+平局双选，钢巴主场赢球难度大。", bold=True)
para("比分：首选1-1，次选0-1，留意1-0。")
para("让球：钢巴账面强但交锋虚，让球方向站不住，看好浦和受让。")
para("大小：2球区间小球为主，大2.5谨慎。")
img("m2_signals.png", "图8｜信号图：交锋克星+赛程分心")

# ===== 互动（本轮轮换） =====
heading("互动｜本轮轮换：数据质疑型 + 战术变数型", level=2)
para("第一场，结论摆这儿：横滨受让+平局第一。市场敢把鹿岛顶到50.6%，我就敢说它看偏了。18个百分点的偏差，是市场在挖坑，还是我的模型太保守？评论区对答案。")
para("第二场，结论：浦和客胜+平局双选。钢巴半场打不开，会不会为亚冠提前收力？留下你的比分，赛后回来验证。")
para("首发评论：两场都站平局的人，这场可能笑得最开心。", italic=True)

heading("配图关键词", level=2)
para("日职联揭幕 横滨水手 鹿岛鹿角 比赛图｜J.League opening Yokohama F. Marinos Kashima Antlers match photos", size=9, color=RGBColor(0x66,0x66,0x66))
para("大阪钢巴 浦和红钻 日职联 比赛图｜Gamba Osaka Urawa Reds J.League match photos", size=9, color=RGBColor(0x66,0x66,0x66))

heading("免责声明", level=2)
para("行文仓促，图片和数据来自官方平台和权威媒体，侵删。", size=9, color=RGBColor(0x88,0x88,0x88))
para("都是我个人图一乐的瞎琢磨，千万别当真，理智看球、远离赌博才是正道。", size=9, color=RGBColor(0x88,0x88,0x88))
para("数据来源：中国体育彩票官方平台 + 新球体育", size=8, color=RGBColor(0xAA,0xAA,0xAA))

doc.save(OUTPUT)
print("saved:", OUTPUT)
